from openpyxl import load_workbook


class XlsxWriter:
    @staticmethod
    def write_task1_scene1_sheet(test_case_id, actual, pass_fail, comments):
        wb = load_workbook("../Test_Reports/TC_Task1_Scenario1.xlsx")
        sheet = wb.active

        for i in range(2, sheet.max_row + 1):
            if sheet.cell(row=i, column=1).value == test_case_id:
                sheet.cell(row=i, column=4, value=actual)  # write to column 'D'
                sheet.cell(row=i, column=5, value=pass_fail)  # write to column 'E'
                sheet.cell(row=i, column=6, value=comments)  # write to column 'F'
                break
        wb.save("../Test_Reports/TC_Task1_Scenario1.xlsx")

    @staticmethod
    def write_task1_scene2_sheet(test_case_id, actual, pass_fail, comments):
        wb = load_workbook("../Test_Reports/TC_Task1_Scenario2.xlsx")
        sheet = wb.active

        for i in range(2, sheet.max_row + 1):
            if sheet.cell(row=i, column=1).value == test_case_id:
                sheet.cell(row=i, column=4, value=actual)  # write to column 'D'
                sheet.cell(row=i, column=5, value=pass_fail)  # write to column 'E'
                sheet.cell(row=i, column=6, value=comments)  # write to column 'F'
                break
        wb.save("../Test_Reports/TC_Task1_Scenario2.xlsx")
