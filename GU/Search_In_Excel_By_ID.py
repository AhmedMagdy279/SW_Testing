import time
import openpyxl


# Excel sheet specs:
# 1) no identifiers in first row
# 2) IDs in first column, Grades second, empty third
# 3) don't forget to change Excel FILE & SHEET variable names
# ************************************************************************************** #
# 4) NOTE, CHECK HOW WE IMPORT ID FROM THE EXCEL, FILES HAS IDs AS str AND OTHERS AS int #
#    IT HAS TO BE str, IF NOT CAST IT TO str(), SAME FOR CELL VALUE TO BE CASTED TO int()#
# ************************************************************************************** #


Target_EXCEL_File = "CW grades.xlsx"
Target_EXCEL_Sheet = "Sheet1"


workbook = openpyxl.load_workbook(Target_EXCEL_File)
worksheet = workbook.get_sheet_by_name(Target_EXCEL_Sheet)


user_name = "Magdy.Eldwaam"
user_pass = "Megz@1962"

driver = webdriver.Firefox()
driver.get("https://sis.gu.edu.eg/psp/ps/EMPLOYEE/SA/c/SSR_ACTIVITY_MGT.SSR_AM_WORKCENTER.GBL?&")
driver.maximize_window()
action = ActionChains(driver)
try:
    if driver.current_url == "https://sis.gu.edu.eg/psp/ps/?cmd=login&errorPg=ckreq&languageCd=ENG":
        driver.find_element(By.XPATH, "//a[text()='Sign in to PeopleSoft']").click()
except:
    print("no entry page")

# entering credentials
driver.find_element(By.ID, "userid").send_keys(user_name)
driver.find_element(By.ID, "pwd").send_keys(user_pass)
driver.find_element(By.XPATH, "//input[@type='submit']").click()

try:
    if driver.find_element(By.ID, "PTNUI_LAND_REC_GROUPLET_LBL$1"):
        driver.find_element(By.ID, "PTNUI_LAND_REC_GROUPLET_LBL$1").click()
        driver.find_element(By.ID, "PTNUI_NB_CNTREC_PTNUI_LINK$0").click()
except:
    print("no entry page")

# checking that login is done
# assert driver.current_url == "https://sis.gu.edu.eg/psp/ps/EMPLOYEE/SA/c/SSR_ACTIVITY_MGT.SSR_AM_WORKCENTER.GBL"


# ##################################### new alternate (commenting this area) ##############################

# time.sleep(2)
# # switch frames
# iframe = driver.find_element(By.ID, "ptalPgltAreaFrame")
# driver.switch_to.frame(iframe)
# # the next part shall be manual to select the course number to add the grades
# # Xpath shall be like this --> //a[text()="LAN 10 (1064)"]
# # entering the course number
# driver.find_element(By.XPATH, "//a[@name='LAN 10 (2774)']").click()
#
# time.sleep(3)
#
# # switch frames
# driver.switch_to.default_content()
# iframe = driver.find_element(By.ID, "ptifrmtgtframe")
# driver.switch_to.frame(iframe)
# driver.find_element(By.ID, "SSR_ACR_RSLT_LINK$IMG$2").click()

# #########################################################################################################

time.sleep(15)
iframe = driver.find_element(By.ID, "ptifrmtgtframe")
driver.switch_to.frame(iframe)
# Locate the table element (adjust the locator if needed)
tables = driver.find_elements(By.XPATH, "//table[@class='PSLEVEL2GRID']")
# Find all table rows starting from the second one (index 1)
all_rows_except_first = tables[1].find_elements(By.XPATH, ".//tr[position() >= 2]")

length = len(all_rows_except_first)


# ((//table[@class='PSLEVEL2GRID'])[2]//tr)[]
# Process each row (excluding the first one)
# Assuming 'length' is defined elsewhere in your code
# creating variable to count IDs found
count = 0
for i in range(0, length):
    # Find the student ID from the web page
    ID_xpath = f"//td/div/span[starts-with(@id, 'EMPLID${i}')]"
    student_id = driver.find_element(By.XPATH, ID_xpath).text

    # loop to check ID in excel sheet
    for row_num in range(1, worksheet.max_row + 1):  # Start from row 1 (no header)
        cell_value1 = str(worksheet.cell(row=row_num, column=1).value)  # Get the value from column A (student ID)
        if cell_value1 == student_id:
            print(f"{student_id} : Found")
            count += 1
            break   # exit out of CW loop if ID is found
print(f"IDs Count : {count}")
driver.close()


